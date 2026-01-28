"""② 見積書（別紙1）の処理"""
import os
from docx import Document
from openpyxl import load_workbook

from .common import clean_formatting, extract_entity_info


def process_estimate(filepath, output_dir, company_name):
    result = {'output_name': '', 'errors': [], 'warnings': [], 'entity_info': None}
    output_name = f'別紙１_{company_name}'
    ext = os.path.splitext(filepath)[1].lower()

    if ext == '.docx':
        doc = Document(filepath)
        clean_formatting(doc)
        result['entity_info'] = extract_entity_info(doc)
        output_name += '.docx'
        doc.save(os.path.join(output_dir, output_name))
    elif ext == '.xlsx':
        wb = load_workbook(filepath)
        output_name += '.xlsx'
        wb.save(os.path.join(output_dir, output_name))
    elif ext == '.pdf':
        import shutil
        output_name += '.pdf'
        shutil.copy2(filepath, os.path.join(output_dir, output_name))
    else:
        result['errors'].append(f'見積書: 未対応のファイル形式です ({ext})')
        return result

    result['output_name'] = output_name
    return result
